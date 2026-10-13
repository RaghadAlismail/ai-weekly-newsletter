"""Arabic newsletter collaborative MVP. Run: python platform/app.py."""
import os, sys, sqlite3, re, json, io, zipfile, datetime, secrets
from pathlib import Path
from urllib.parse import urlsplit
from flask import Flask, request, redirect, url_for, render_template, flash, abort, send_file, jsonify, session, send_from_directory
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix
from storage import database, initialize
from security import init_google, verified_identity, csrf_token, enforce_csrf, authorized_people

BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE))
DATA=Path(os.getenv('NEWSLETTER_DATA_DIR',str(BASE/'platform'/'data')))
DATA.mkdir(parents=True,exist_ok=True)
app=Flask(__name__,template_folder=str(BASE),static_folder=None)
@app.route('/static/<path:filename>',endpoint='static')
def public_stylesheet(filename):
    if filename!='style.css':abort(404)
    return send_from_directory(BASE,filename)
if os.getenv('APP_ENV')=='production' and not os.getenv('FLASK_SECRET_KEY'):
    raise RuntimeError('FLASK_SECRET_KEY is required in production')
app.secret_key=os.getenv('FLASK_SECRET_KEY') or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_SECURE=os.getenv('APP_ENV')=='production',SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=8))
app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=1)
oauth=init_google(app)
app.before_request(enforce_csrf)
@app.after_request
def secure_headers(resp):
    resp.headers['X-Content-Type-Options']='nosniff'
    resp.headers['Referrer-Policy']='strict-origin-when-cross-origin'
    resp.headers['X-Frame-Options']='DENY'
    resp.headers['Content-Security-Policy']="default-src 'self'; img-src 'self' https: data:; style-src 'self'; form-action 'self'; base-uri 'none'; frame-ancestors 'none'"
    if os.getenv('APP_ENV')=='production':resp.headers['Strict-Transport-Security']='max-age=31536000; includeSubDomains'
    return resp

app.config['MAX_CONTENT_LENGTH']=6*1024*1024
CATEGORIES=['حكومات','شركات','دراسات','جامعات','منظمات']
STAGES=['content','language','design','final']
STAGE_AR={'content':'مراجعة المحتوى','language':'التدقيق اللغوي','design':'مراجعة التصميم','final':'الاعتماد النهائي'}
ROLE_PERMISSIONS={'coordinator':{'content','language','design'},'content_reviewer':{'content'},'language_reviewer':{'language'},'designer':{'design'},'approver':{'final'},'admin':set(STAGES)}

def db():return database(DATA)
initialize(DATA)

def now(): return datetime.datetime.now(datetime.timezone.utc).isoformat()
def actor(): return (session.get('email') or '').strip()
def role():
    email=actor().casefold()
    return authorized_people().get(email,'') if email else ''

def allow_roles(*roles):
    require_login()
    if role() not in roles:abort(403,'This action requires an assigned editorial role.')

def require_login():
    if not actor() or not role():abort(403,'Sign in with an authorized Google account.')
def writable(i):
    if i['status']=='approved': abort(409,'Approved master is locked. Create a new issue or unlock through a controlled revision process.')
def issue(iid):
    with db() as c: r=c.execute('SELECT * FROM issues WHERE id=?',(iid,)).fetchone()
    if not r: abort(404)
    return r
def article(aid):
    with db() as c: r=c.execute('SELECT * FROM articles WHERE id=?',(aid,)).fetchone()
    if not r: abort(404)
    return r
def audit(c,iid,aid,action,payload=None):
    c.execute('INSERT INTO audit(issue_id,article_id,action,actor,payload) VALUES(?,?,?,?,?)',(iid,aid,action,actor() or 'system',json.dumps(payload or {},ensure_ascii=False)))

def safeurl(s):
    s=s.strip()
    p=urlsplit(s)
    if p.scheme not in ('https','http') or not p.hostname or p.username or p.password or p.port not in (None,80,443): return False
    if p.hostname.lower() in ('localhost','127.0.0.1','0.0.0.0') or p.hostname.lower().endswith(('.local','.internal')): return False
    import ipaddress
    try:
        a=ipaddress.ip_address(p.hostname)
        if not a.is_global:return False
    except ValueError:pass
    return True

def verify_access(aid):
    a=article(aid); i=issue(a['issue_id']);require_login(); writable(i);
    from workflow import check_edit
    check_edit(i,actor(),role());return a,i

@app.route('/login')
def login():
    if actor() and role():return redirect(url_for('home'))
    return render_template('login.html')

@app.route('/auth/google')
def google_start():
    client=oauth.create_client('google')
    if client is None:abort(503,'Google OAuth is not configured')
    return client.authorize_redirect(url_for('google_callback',_external=True))

@app.route('/connect/google-publishing')
def connect_google_publishing():
    require_login()
    session['publishing_connect_email']=actor().casefold()
    client=oauth.create_client('google')
    if client is None:abort(503,'Google authentication is not configured')
    from google_publish import SCOPES
    return client.authorize_redirect(url_for('google_callback',_external=True),scope='openid email profile '+' '.join(SCOPES),access_type='offline',prompt='consent')

@app.route('/auth/callback')
def google_callback():
    try:
        token=oauth.google.authorize_access_token()
        info=token.get('userinfo')
        if not info:abort(403,'Google OpenID identity is missing')
        identity=verified_identity(info)
    except Exception:
        app.logger.exception('OAuth sign-in failed')
        abort(403,'Google sign-in failed')
    if not identity:abort(403,'Google account is not authorized for this workspace.')
    connect_email=session.get('publishing_connect_email')
    if connect_email:
        if connect_email!=identity['email']:abort(403,'Connect publishing with the same Google account used for the workspace.')
        from google_publish import store_google_token,SCOPES
        if not set(SCOPES).issubset(set(str(token.get('scope','')).split())):abort(403,'Google document permissions were not granted.')
        store_google_token(identity['email'],token)
    session.clear()
    session.update(email=identity['email'], display_name=identity['name'])
    session.permanent=True
    return redirect(url_for('home'))

@app.route('/logout',methods=['POST'])
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/healthz')
def healthz():
    try:
        with db() as c:c.execute('SELECT 1')
    except Exception:return jsonify({'status':'unhealthy'}),503
    return jsonify({'status':'ok'})

@app.route('/workspace')
def home():
    require_login()
    with db() as c: issues=c.execute('SELECT i.*, (SELECT COUNT(*) FROM articles WHERE issue_id=i.id) n,(SELECT COUNT(*) FROM articles WHERE issue_id=i.id AND headline<>\'\') ready FROM issues i ORDER BY i.id DESC').fetchall()
    return render_template('home.html',issues=issues)
# Only explicitly released, sanitized snapshots are exposed on public routes.
# Approving the Arabic master does NOT make it public.
def public_payload(approved):
    return {'issue':approved['issue'],'date':approved['date'],
       'articles':[{'position':a['position'],'headline':a['headline'],
       'body':a['body'],'category':a['category'],'url':a['url'],
       'image_local':a.get('image_local','') if a.get('image_license_status')=='approved' else ''}
       for a in approved['articles']]}

@app.route('/')
def public_home():
    with db() as c:
        releases=c.execute('SELECT r.issue_id,r.public_snapshot,r.released_at FROM public_releases r WHERE r.revoked_at IS NULL ORDER BY r.issue_id DESC LIMIT 100').fetchall()
    listing=[{'id':r['issue_id'],'number':json.loads(r['public_snapshot'])['issue'],
              'date':json.loads(r['public_snapshot'])['date']} for r in releases]
    return render_template('public_home.html',listing=listing)

@app.route('/public/issues/<int:iid>')
def public_issue(iid):
    with db() as c:
        release=c.execute('SELECT public_snapshot,released_at FROM public_releases WHERE issue_id=? AND revoked_at IS NULL',(iid,)).fetchone()
    if not release:abort(404)
    payload=json.loads(release['public_snapshot'])
    return render_template('public_issue.html',edition=payload,iid=iid)

@app.route('/public/issues/<int:iid>/images/<int:position>')
def public_image(iid,position):
    with db() as c:
        release=c.execute('SELECT public_snapshot FROM public_releases WHERE issue_id=? AND revoked_at IS NULL',(iid,)).fetchone()
    if not release:abort(404)
    articles=json.loads(release['public_snapshot'])['articles']
    match=next((a for a in articles if a['position']==position),None)
    if not match or not match.get('image_local'):abort(404)
    name=Path(match['image_local']).name
    if name!=match['image_local']:abort(404)
    return send_from_directory(DATA/'images',name,max_age=3600)

@app.route('/issues/<int:iid>/release',methods=['POST'])
def release_issue(iid):
    allow_roles('approver','admin')
    i=issue(iid)
    if i['status']!='approved' or not i['approved_snapshot']:abort(409,'Only an approved Arabic issue can be released.')
    sanitized=public_payload(json.loads(i['approved_snapshot']))
    with db() as c:
        c.execute('INSERT INTO public_releases(issue_id,public_snapshot,released_by,revoked_at) VALUES(?,?,?,NULL) ON CONFLICT(issue_id) DO UPDATE SET public_snapshot=excluded.public_snapshot,released_by=excluded.released_by,revoked_at=NULL,released_at=CURRENT_TIMESTAMP',
                  (iid,json.dumps(sanitized,ensure_ascii=False),actor()))
        audit(c,iid,None,'public_released')
    flash('The approved issue is now publicly readable.','success')
    return redirect(url_for('view_issue',iid=iid))

@app.route('/issues/<int:iid>/unrelease',methods=['POST'])
def unrelease_issue(iid):
    allow_roles('approver','admin')
    issue(iid)
    with db() as c:
        c.execute('UPDATE public_releases SET revoked_at=CURRENT_TIMESTAMP WHERE issue_id=?',(iid,))
        audit(c,iid,None,'public_unreleased')
    flash('The public issue has been withdrawn.','success')
    return redirect(url_for('view_issue',iid=iid))

@app.route('/issues/new',methods=['GET','POST'])
def new_issue():
    allow_roles('coordinator','admin')
    if request.method=='POST':
        with db() as c:
            next_number=c.execute('SELECT COALESCE(MAX(number),258)+1 AS n FROM issues').fetchone()['n']
        try:
            num=int(request.form.get('number','').strip() or next_number)
            date=request.form.get('date','').strip() or datetime.date.today().isoformat()
            if num<1:raise ValueError()
            datetime.date.fromisoformat(date)
        except ValueError:
            flash('Use a positive issue number and a valid publication date, or leave them blank.','error');return redirect(request.url)
        entries=[]
        for j in range(1,11):
            u=request.form.get('url_'+str(j),'').strip()
            if not u:continue
            category=request.form.get('category_'+str(j),'').strip()
            img=request.form.get('image_'+str(j),'').strip()
            if not safeurl(u) or (category and category not in CATEGORIES) or (img and not safeurl(img)):
                flash('Use valid public URLs and an available category.','error');return redirect(request.url)
            entries.append((u,request.form.get('content_'+str(j),'').strip(),request.form.get('title_'+str(j),'').strip(),category,img))
        if len(entries)!=10 or len({x[0] for x in entries})!=10:
            flash('Enter exactly ten distinct article URLs. All other fields are optional.','error');return redirect(request.url)
        with db() as c:
            try:
                cur=c.execute('INSERT INTO issues(number,pub_date,owner_email) VALUES(?,?,?)'+(' RETURNING id' if os.getenv('DATABASE_URL') else ''),(num,date,actor().casefold()));iid=cur.fetchone()['id'] if os.getenv('DATABASE_URL') else cur.lastrowid
            except (sqlite3.IntegrityError, __import__('psycopg').IntegrityError):flash('Issue number already exists','error');return redirect(request.url)
            for j,(u,content,title,category,img) in enumerate(entries,1):
                c.execute('INSERT INTO articles(issue_id,position,url,supplied_content,source_title,category,image_url) VALUES(?,?,?,?,?,?,?)',(iid,j,u,content,title,category,img))
            audit(c,iid,None,'issue_created',{'urls':len(entries)})
        return app.view_functions['generate_draft'](iid)
    return render_template('new.html',date=datetime.date.today().isoformat())
@app.route('/issues/<int:iid>')
def view_issue(iid):
    require_login();i=issue(iid)
    with db() as c:
        aa=c.execute('SELECT * FROM articles WHERE issue_id=? ORDER BY position',(iid,)).fetchall()
        reviews=c.execute('SELECT * FROM reviews WHERE issue_id=? ORDER BY created_at DESC,id DESC',(iid,)).fetchall()
        pubs=c.execute('SELECT * FROM publications WHERE issue_id=? ORDER BY id DESC',(iid,)).fetchall()
    return render_template('issue.html',i=i,aa=aa,reviews=reviews,pubs=pubs,stages=STAGE_AR)
@app.route('/articles/<int:aid>')
def edit_article(aid):
    require_login();a=article(aid);i=issue(a['issue_id'])
    with db() as c:reviews=c.execute('SELECT * FROM reviews WHERE article_id=? ORDER BY id DESC',(aid,)).fetchall()
    return render_template('article.html',a=a,i=i,categories=CATEGORIES,reviews=reviews,stages=STAGE_AR)
def lock_article_edit(c,i):
    locked=c.execute('SELECT * FROM issues WHERE id=?'+(' FOR UPDATE' if os.getenv('DATABASE_URL') else ''),(i['id'],)).fetchone()
    if locked['status']=='approved' or locked['workflow_stage']!=i['workflow_stage'] or locked['content_revision']!=i['content_revision']:abort(409,'The newsletter changed. Refresh before saving.')

@app.route('/articles/<int:aid>/save',methods=['POST'])
def save(aid):
    allow_roles('coordinator','content_reviewer','language_reviewer','designer','approver','admin')
    a,i=verify_access(aid)
    updates={k:request.form.get(k,'').strip() for k in ('headline','body','category','image_url','image_asset_id','image_license_status','supplied_content')}
    if updates['category'] not in CATEGORIES:abort(400)
    if updates['image_license_status'] not in ('unverified','approved','not_required'):abort(400)
    if updates['image_url'] and not safeurl(updates['image_url']):abort(400)
    new_image=None
    if updates['image_url'] and updates['image_url']!=a['image_url']:
        from source_fetch import fetch_image
        try:new_image=fetch_image(updates['image_url'])
        except Exception:flash('The image URL could not be loaded. Use a valid image URL or upload an image.','error');return redirect(url_for('edit_article',aid=aid))
    with db() as c:
        lock_article_edit(c,i)
        c.execute('UPDATE articles SET headline=?,body=?,category=?,image_url=?,image_asset_id=?,image_license_status=?,supplied_content=?,revision=revision+1 WHERE id=?',(*updates.values(),aid))
        audit(c,i['id'],aid,'article_saved',{'new_revision':a['revision']+1})
        if new_image:
            import base64
            c.execute('INSERT INTO article_media(article_id,payload) VALUES(?,?) ON CONFLICT(article_id) DO UPDATE SET payload=excluded.payload',(aid,base64.b64encode(new_image).decode()))
        workflow_helpers['touch'](i['id'],'Article content edited',any(updates[k]!=a[k] for k in ('headline','body','category','supplied_content')),c)
    flash('Saved. Previous approvals are no longer valid for this article.','success')
    return redirect(url_for('edit_article',aid=aid))
@app.route('/articles/<int:aid>/upload-image',methods=['POST'])
def upload_image(aid):
    allow_roles('coordinator','content_reviewer','language_reviewer','designer','approver','admin')
    a,i=verify_access(aid)
    f=request.files.get('image')
    if not f or not f.filename or Path(f.filename).suffix.lower() not in ('.png','.jpg','.jpeg','.webp'):
        flash('Upload a JPEG, PNG or WebP image.','error');return redirect(url_for('edit_article',aid=aid))
    import PIL.Image
    buf=f.read(5*1024*1024+1)
    if len(buf)>5*1024*1024:abort(413)
    try:
        im=PIL.Image.open(io.BytesIO(buf)); im.verify()
    except Exception:abort(400,'Not a valid image')
    name=f'article_{aid}_{secrets.token_hex(6)}'+Path(f.filename).suffix.lower()
    folder=DATA/'images';folder.mkdir(exist_ok=True)
    (folder/name).write_bytes(buf)
    with db() as c:
        lock_article_edit(c,i)
        c.execute("UPDATE articles SET image_local=?,image_license_status='unverified',revision=revision+1 WHERE id=?",(name,aid))
        audit(c,i['id'],aid,'image_uploaded',{'filename':name})
        import base64
        im=PIL.Image.open(io.BytesIO(buf)).convert('RGB');media=io.BytesIO();im.save(media,'JPEG',quality=94)
        c.execute('INSERT INTO article_media(article_id,payload) VALUES(?,?) ON CONFLICT(article_id) DO UPDATE SET payload=excluded.payload',(aid,base64.b64encode(media.getvalue()).decode()))
        workflow_helpers['touch'](i['id'],'Image updated',False,c)
    flash('Image uploaded. Confirm rights and save before design approval.','success')
    return redirect(url_for('edit_article',aid=aid))

@app.route('/articles/<int:aid>/extract',methods=['POST'])
def extract(aid):
    allow_roles('coordinator','content_reviewer','admin')
    a,i=verify_access(aid)
    if a['supplied_content'].strip():
        title=a['source_title'] or 'User-supplied article';text=a['supplied_content'];result={'image_url':a['image_url'],'manual_review_required':False};status='supplied'
    else:
        from source_fetch import fetch_article
        try:
            result=fetch_article(a['url']);title=result.get('title') or '';text=result.get('text') or '';status='ready' if len(text)>=200 and not result.get('manual_input_required') else 'needs_input'
        except Exception as ex:title='';text='';result={'error':str(ex)};status='needs_input'
    with db() as c:
        c.execute('UPDATE articles SET source_title=?,source_text=?,extract_status=?,extract_notes=?,image_url=CASE WHEN image_url=\'\' THEN ? ELSE image_url END,revision=revision+1 WHERE id=?',(title,text,status,json.dumps(result.get('fetch_warnings',[]) or [result.get('error','')],ensure_ascii=False),result.get('image_url','') or '',aid));audit(c,i['id'],aid,'article_extracted',{'status':status})
    flash('Source extraction: '+status,'success' if status!='needs_input' else 'error');return redirect(url_for('edit_article',aid=aid))
@app.route('/articles/<int:aid>/generate',methods=['POST'])
def generate(aid):
    allow_roles('coordinator','content_reviewer','admin')
    a,i=verify_access(aid)
    content=a['source_text'] or a['supplied_content']
    if len(content.strip())<200:flash('Provide or extract sufficient source content before generating.','error');return redirect(url_for('edit_article',aid=aid))
    if not os.getenv('OPENAI_API_KEY'):
        flash('OPENAI_API_KEY is not configured. You can edit Arabic content manually.','error');return redirect(url_for('edit_article',aid=aid))
    try:
        from write_summary import write_summary
        # Map actual approved category to legacy writer only for prompt; replace with issue-specific instructions afterwards.
        category_en={'حكومات':'Governments','شركات':'Companies','جامعات':'Universities','منظمات':'Organizations','دراسات':'Studies'}[a['category']]
        result=write_summary(a['source_title'] or a['url'],content,category_en)
        headline=result.get('title_ar','').strip();body=result.get('body_ar','').strip()
        if not headline or not body:raise ValueError('Model output missing Arabic fields')
        with db() as c:
            c.execute('UPDATE articles SET headline=?,body=?,revision=revision+1 WHERE id=?',(headline,body,aid));audit(c,i['id'],aid,'ai_generated')
        workflow_helpers['touch'](i['id'],'Article regenerated',True)
        flash('Draft generated. Human review is required.','success')
    except Exception as exc:flash('Generation failed: '+str(exc)[:220],'error')
    return redirect(url_for('edit_article',aid=aid))
@app.route('/articles/<int:aid>/review',methods=['POST'])
def review(aid):
    abort(409,'Use the whole-newsletter review decision on the issue dashboard.')
    allow_roles('coordinator','content_reviewer','language_reviewer','designer','approver','admin')
    a,i=verify_access(aid)
    stage=request.form.get('stage','');decision=request.form.get('decision','')
    if stage not in STAGES[:-1] or stage not in ROLE_PERMISSIONS.get(role(),set()) or decision not in ('approved','changes_requested'):abort(403)
    if decision=='approved' and (not a['headline'].strip() or not a['body'].strip() or a['extract_status'] not in ('ready','supplied')):
        flash('Extract source and complete headline/body before approval.','error');return redirect(url_for('edit_article',aid=aid))
    if stage=='design' and decision=='approved' and (a['image_license_status']!='approved' or not (a['image_local'] or a['image_url'])):flash('Confirm image rights before design approval.','error');return redirect(url_for('edit_article',aid=aid))
    with db() as c:
        c.execute('INSERT INTO reviews(issue_id,article_id,stage,revision,reviewer,decision,note) VALUES(?,?,?,?,?,?,?)',(i['id'],aid,stage,a['revision'],actor(),decision,request.form.get('note','')[:1500]));audit(c,i['id'],aid,'review',{'stage':stage,'decision':decision})
    return redirect(url_for('edit_article',aid=aid))
def latest_reviews(c,iid):
    aa=c.execute('SELECT * FROM articles WHERE issue_id=? ORDER BY position',(iid,)).fetchall()
    issues=[]
    for a in aa:
        for stage in STAGES[:-1]:
            rr=c.execute('SELECT decision FROM reviews WHERE article_id=? AND stage=? AND revision=? ORDER BY id DESC LIMIT 1',(a['id'],stage,a['revision'])).fetchone()
            if not rr or rr['decision']!='approved':issues.append(f'Article {a["position"]}: {STAGE_AR[stage]} pending')
    return aa,issues
@app.route('/issues/<int:iid>/approve',methods=['POST'])
def approve(iid):
    abort(409,'Use the sequential final reviewer approval on the issue dashboard.')
    require_login();i=issue(iid);writable(i)
    if role() not in ('approver','admin'):abort(403)
    with db() as c:
        aa,missing=latest_reviews(c,iid)
        if not 1<=len(aa)<=10 or missing:
            flash('Cannot approve: '+('; '.join(missing[:4]) or 'One to ten articles required'),'error');return redirect(url_for('view_issue',iid=iid))
        snapshot={'issue':i['number'],'date':i['pub_date'],'articles':[{k:a[k] for k in ('position','url','headline','body','category','image_url','image_local','image_asset_id','image_license_status','revision')} for a in aa]}
        snap=json.dumps(snapshot,ensure_ascii=False)
        c.execute('UPDATE issues SET status=\'approved\',approved_at=?,approved_snapshot=? WHERE id=?',(now(),snap,iid));c.execute('INSERT INTO reviews(issue_id,article_id,stage,revision,reviewer,decision) VALUES(?,NULL,\'final\',0,?,\'approved\')',(iid,actor()));audit(c,iid,None,'master_approved')
    flash('Arabic master approved and locked. Short version is now available.','success');return redirect(url_for('view_issue',iid=iid))
@app.route('/issues/<int:iid>/short.docx')
def short_doc(iid):
    require_login();i=issue(iid)
    if i['status']!='approved':abort(403,'Short version is generated only after Arabic approval.')
    out=build_short_doc(json.loads(i['approved_snapshot']))
    return send_file(out,mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',as_attachment=True,download_name=f'newsletter_{i["number"]}_short.docx')

def build_short_doc(data):
    from docx import Document
    d=Document(); d.add_heading(f'نشرة الذكاء الاصطناعي {data["issue"]}',0)
    for a in data['articles']:
        d.add_heading(f'الشريحة {a["position"]+1}',level=1)
        d.add_paragraph('الصورة:')
        if a.get('image_local'):
            img=DATA/'images'/Path(a['image_local']).name
            if img.is_file():
                try:d.add_picture(str(img),width=__import__('docx').shared.Inches(5.6))
                except Exception:d.add_paragraph('Image insertion failed; verify file format')
        else:d.add_paragraph(a['image_url'] or 'Image pending')
        for label,val in [('العنوان',a['headline']),('التصنيف',a['category']),('الرابط',a['url'])]:
            d.add_paragraph(f'{label}: {val}')
    out=io.BytesIO();d.save(out);out.seek(0)
    return out

@app.route('/issues/<int:iid>/publish-short',methods=['POST'])
def publish_short(iid):
    require_login();i=issue(iid)
    if i['status']!='approved' or role() not in ('approver','admin'):abort(403)
    data=json.loads(i['approved_snapshot'])
    try:
        from google_publish import create_google_short_doc
        url,id_=create_google_short_doc(build_short_doc(data).getvalue(),data)
        with db() as c:
            c.execute('INSERT INTO publications(issue_id,kind,external_id,source_snapshot,status) VALUES(?,?,?,?,?)',(iid,'arabic_short_doc',id_,i['approved_snapshot'],'created'))
            audit(c,iid,None,'published_short_doc',{'document_id':id_})
        flash('Native Google Doc created: '+url,'success')
    except Exception as ex:flash('Google Docs setup required: '+str(ex)[:260],'error')
    return redirect(url_for('view_issue',iid=iid))
@app.route('/issues/<int:iid>/manifest.json')
def manifest(iid):
    require_login();i=issue(iid)
    if i['status']!='approved':abort(403)
    return jsonify(json.loads(i['approved_snapshot']))
@app.route('/issues/<int:iid>/publish-slides',methods=['POST'])
def publish_slides(iid):
    require_login();i=issue(iid)
    if i['status']!='approved':abort(403)
    if role() not in ('approver','admin'):abort(403)
    try:
        from google_publish import create_google_slides
        from deck_builder import build_deck
        data=json.loads(i['approved_snapshot'])
        url,id_=create_google_slides(data,build_deck(data,workflow_helpers['media'](iid)))
        with db() as c:c.execute('INSERT INTO publications(issue_id,kind,external_id,source_snapshot,status) VALUES(?,?,?,?,?)',(iid,'arabic_slides',id_,i['approved_snapshot'],'created'));audit(c,iid,None,'published_slides',{'presentation_id':id_})
        flash('Native Google Slides created: '+url,'success')
    except Exception as ex:flash('Google Slides setup required: '+str(ex)[:260],'error')
    return redirect(url_for('view_issue',iid=iid))
@app.context_processor
def context():return {'current_user':session.get('display_name',actor()),'current_role':role(),'STAGE_AR':STAGE_AR,'csrf_token':csrf_token}
from workflow import register_workflow
workflow_helpers={'db':db,'issue':issue,'actor':actor,'role':role,'require_login':require_login,'authorized_people':authorized_people,'audit':audit,'DATA':DATA,'CATEGORIES':CATEGORIES,'now':now}
register_workflow(app,workflow_helpers)
if __name__=='__main__':app.run(host='127.0.0.1',port=int(os.getenv('PORT','5055')),debug=False)
